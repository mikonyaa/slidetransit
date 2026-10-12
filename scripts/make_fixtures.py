"""Create our own PPTX decks and genuinely export them using an isolated profile.

No PDF is drawn here. All PDFs in this corpus are produced by soffice.
The font directory must contain only fonts whose embedding license permits it.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import platform
import shutil
import subprocess
import tempfile
from datetime import UTC, datetime
from pathlib import Path
from xml.sax.saxutils import escape

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.text import MSO_AUTO_SIZE
from pptx.util import Pt


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def textbox(slide, name, text, x, y, width, height, size=20, font="Noto Sans"):
    shape = slide.shapes.add_textbox(Pt(x), Pt(y), Pt(width), Pt(height))
    shape.name = name
    frame = shape.text_frame
    frame.margin_left = frame.margin_right = frame.margin_top = frame.margin_bottom = 0
    frame.auto_size = MSO_AUTO_SIZE.NONE
    frame.word_wrap = True
    frame.text = text
    for para in frame.paragraphs:
        para.font.name = font
        para.font.size = Pt(size)
        para.font.color.rgb = RGBColor.from_string("172926")
        for run in para.runs:
            run.font.name = font
            run.font.size = Pt(size)
            # Explicit East Asian font avoids an undocumented system fallback.
            from pptx.oxml.xmlchemy import OxmlElement

            ea = OxmlElement("a:ea")
            ea.set("typeface", "Noto Sans CJK SC")
            run._r.get_or_add_rPr().append(ea)
    return shape


def make_deck(path: Path, variant: str, image: Path) -> None:
    prs = Presentation()
    prs.core_properties.author = "SlideTransit fixtures"
    prs.core_properties.last_modified_by = "SlideTransit fixtures"
    prs.core_properties.title = "Delivery study - own documented test variants"
    prs.slide_width, prs.slide_height = Pt(720), Pt(405)
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    textbox(slide, "Kicker", "SLIDETRANSIT / OWN TEST DECK", 42, 30, 630, 22, 12)
    title = textbox(slide, "Delivery title", "Delivery without surprises", 42, 82, 630, 65, 34)
    if variant == "clipping":
        # Deliberately place the bottom of a real title outside the slide.
        # Text can survive extraction while the rendered page visibly clips it.
        title.top = Pt(382)
    if variant == "wrap":
        title.width, title.height = Pt(280), Pt(104)
    if variant == "displacement":
        title.left = Pt(70)
    if variant != "deleted-text":
        textbox(
            slide,
            "Description",
            "Compare the exports. Keep the evidence.",
            42,
            207,
            600,
            34,
            21,
            "Liberation Serif" if variant == "font" else "Noto Sans",
        )
    textbox(
        slide,
        "Metric",
        "Conversion 12.6%" if variant == "symbol" else "Conversion 12.5%",
        42,
        275,
        410,
        30,
        20,
    )
    textbox(
        slide,
        "CJK",
        "東京 中文 演示" if variant != "cjk" else "東京 中文 演示版本",
        42,
        330,
        410,
        30,
        20,
    )
    if variant != "deleted-image":
        slide.shapes.add_picture(str(image), Pt(556), Pt(258), Pt(110), Pt(80))
    second = prs.slides.add_slide(prs.slide_layouts[6])
    textbox(second, "Title", "Evidence survives delivery", 42, 42, 640, 48, 30)
    textbox(second, "Review A", "Review", 42, 136, 180, 35, 22)
    textbox(second, "Review B", "Review", 372 if variant != "duplicate" else 400, 136, 180, 35, 22)
    textbox(second, "Local", "All input bytes stay local.", 42, 240, 610, 35, 24)
    if variant == "pages":
        third = prs.slides.add_slide(prs.slide_layouts[6])
        textbox(
            third, "Extra", "An inserted page makes correspondence uncertain.", 42, 60, 620, 65, 25
        )
    if variant == "aspect":
        prs.slide_width = Pt(760)
    prs.save(path)


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out", type=Path, default=Path("tests/fixtures/real"))
    parser.add_argument("--font-dir", type=Path, required=True)
    parser.add_argument("--soffice", default="soffice")
    args = parser.parse_args()
    if args.out.exists():
        parser.error("fixture output already exists; use a new directory")
    args.out.mkdir(parents=True)
    fonts = args.font_dir.resolve()
    version = subprocess.check_output([args.soffice, "--version"], text=True, timeout=30).strip()
    environment = {
        "os": platform.system(),
        "osVersion": platform.mac_ver()[0] or platform.release(),
        "architecture": platform.machine(),
        "python": platform.python_version(),
    }
    exports = {}
    with tempfile.TemporaryDirectory(prefix="slidetransit-export-") as temporary:
        work = Path(temporary)
        config = work / "fonts.conf"
        config.write_text(
            f"<fontconfig><dir>{escape(str(fonts))}</dir><cachedir>{escape(str(work / 'cache'))}</cachedir></fontconfig>"
        )
        env = dict(os.environ, FONTCONFIG_FILE=str(config), XDG_CACHE_HOME=str(work / "cache"))
        image = work / "study.png"
        canvas = Image.new("RGB", (440, 320), "#edf0e8")
        draw = ImageDraw.Draw(canvas)
        for index, height in enumerate([80, 145, 200, 250]):
            draw.rectangle((40 + index * 95, 280 - height, 99 + index * 95, 280), fill="#bc5135")
        canvas.save(image)
        noise_image = work / "noise.png"
        canvas.putpixel((15, 15), (231, 234, 226))
        canvas.save(noise_image)
        variants = [
            "baseline",
            "repeat",
            "clipping",
            "deleted-text",
            "deleted-image",
            "wrap",
            "font",
            "symbol",
            "cjk",
            "displacement",
            "duplicate",
            "pages",
            "aspect",
            "noise",
            "fixed",
        ]
        for variant in variants:
            source = args.out / (variant + ".pptx")
            if variant in ["repeat", "fixed"]:
                shutil.copyfile(args.out / "baseline.pptx", source)
            else:
                make_deck(source, variant, noise_image if variant == "noise" else image)
            profile = work / ("profile-" + variant)
            command = [
                args.soffice,
                "-env:UserInstallation=" + profile.as_uri(),
                "--headless",
                "--convert-to",
                "pdf:impress_pdf_Export",
                "--outdir",
                str(args.out.resolve()),
                str(source.resolve()),
            ]
            result = subprocess.run(
                command, env=env, capture_output=True, text=True, timeout=90, check=True
            )
            pdf = source.with_suffix(".pdf")
            if not pdf.is_file():
                raise RuntimeError("export did not create PDF: " + result.stderr)
            exports[variant] = {
                "path": pdf.name,
                "renderer": {"name": "LibreOfficeDev", "version": version},
                "source": {"path": source.name, "sha256": sha(source)},
                "sha256": sha(pdf),
                "provenanceAssurance": "captured-by-fixture-script",
                "exportOptions": {"filter": "impress_pdf_Export", "headless": True},
                "environment": environment,
                "capturedAt": datetime.now(UTC).isoformat(),
                "command": [
                    "soffice",
                    "-env:UserInstallation=file://<isolated-task-profile>",
                    "--headless",
                    "--convert-to",
                    "pdf:impress_pdf_Export",
                    "--outdir",
                    ".",
                    source.name,
                ],
            }
            print(f"exported {variant}: {pdf.stat().st_size} bytes", flush=True)
    for variant in variants[1:]:
        manifest = {
            "schemaVersion": "1.0",
            "presentation": {"path": "baseline.pptx", "sha256": sha(args.out / "baseline.pptx")},
            "mode": "declared-same-source"
            if variant in ["repeat", "fixed"]
            else "documented-test-variant",
            "reference": exports["baseline"],
            "candidate": exports[variant],
            "scope": {
                "slideRange": [1, 2],
                "pageOrder": "source-order",
                "kind": "slides",
                "includeHidden": False,
            },
            "limitations": [
                "Own test variant; no cross-exporter regression claim.",
                "Fixture exporter is an explicitly identified development build.",
            ],
        }
        (args.out / (variant + ".json")).write_text(
            json.dumps(manifest, indent=2, ensure_ascii=False) + "\n"
        )
    (args.out / "provenance.json").write_text(
        json.dumps(exports, indent=2, ensure_ascii=False) + "\n"
    )


if __name__ == "__main__":
    main()
